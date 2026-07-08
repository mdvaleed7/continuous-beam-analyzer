#!/usr/bin/env python3
"""
Structural Engineering Handbook — Markdown → DOCX converter.

Reads the numbered chapter markdown files in this directory (00-...md through
08-...md) and produces a single, styled Word document.

Supported markdown subset:
  #, ##, ###, #### headings
  - bullets (one nesting level with two-space indent)
  1. numbered lists
  | tables |
  **bold**, *italic*, `code` inline
  > blockquotes
  --- horizontal rules (rendered as spacing)

Usage:  python3 build_docx.py [output.docx]
"""

import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)      # deep navy
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)     # word-blue
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_BG = "1F3A5F"
ALT_ROW_BG = "EDF2F8"


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_runs_with_inline_md(paragraph, text, base_size=None, base_color=None):
    """Parse **bold**, *italic*, `code` inline markdown into runs."""
    token_re = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    pos = 0
    for m in token_re.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            _style_run(r, base_size, base_color)
        tok = m.group(0)
        if tok.startswith("**"):
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
        else:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
        _style_run(r, base_size, base_color)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        _style_run(r, base_size, base_color)


def _style_run(run, size, color):
    if size:
        run.font.size = size
    if color:
        run.font.color.rgb = color


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, bold, before, after in [
        ("Heading 1", 20, ACCENT, True, 18, 10),
        ("Heading 2", 15, ACCENT2, True, 14, 6),
        ("Heading 3", 12.5, ACCENT, True, 10, 4),
        ("Heading 4", 11, GREY, True, 8, 3),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = bold
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_table(doc, header, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for j, htxt in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        add_runs_with_inline_md(p, htxt)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, HEADER_BG)

    for i, row in enumerate(rows, start=1):
        for j in range(len(header)):
            txt = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs_with_inline_md(p, txt, base_size=Pt(9.5))
            if i % 2 == 0:
                set_cell_bg(cell, ALT_ROW_BG)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def process_markdown(doc, md_text, first_heading_page_break=False):
    lines = md_text.split("\n")
    i = 0
    first_h1_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # table detection
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s\-:|]+\|\s*$", lines[i + 1]
        ):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, header, rows)
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                if first_heading_page_break and not first_h1_done:
                    doc.add_page_break()
                first_h1_done = True
            h = doc.add_heading("", level=level)
            add_runs_with_inline_md(h, text)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            add_runs_with_inline_md(p, stripped.lstrip("> ").strip(),
                                    base_color=GREY)
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        mb = re.match(r"^(\s*)-\s+(.*)$", line)
        if mb:
            indent = len(mb.group(1))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            # merge continuation lines
            text = mb.group(2)
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if nxt.strip() and not re.match(r"^\s*([-#>|]|\d+\.)", nxt) and (
                    len(nxt) - len(nxt.lstrip()) >= indent + 2
                ):
                    text += " " + nxt.strip()
                    i += 1
                else:
                    break
            p = doc.add_paragraph(style=style)
            add_runs_with_inline_md(p, text)
            continue

        mn = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if mn:
            text = mn.group(2)
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if nxt.strip() and not re.match(r"^\s*([-#>|]|\d+\.)", nxt) and (
                    len(nxt) - len(nxt.lstrip()) >= 2
                ):
                    text += " " + nxt.strip()
                    i += 1
                else:
                    break
            p = doc.add_paragraph(style="List Number")
            add_runs_with_inline_md(p, text)
            continue

        # normal paragraph — merge soft-wrapped lines
        text = stripped
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if nxt and not re.match(r"^([-#>|]|\d+\.|---)", nxt) and not nxt.startswith("|"):
                text += " " + nxt
                i += 1
            else:
                break
        p = doc.add_paragraph()
        add_runs_with_inline_md(p, text)


def build_title_page(doc):
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("STRUCTURAL ENGINEERING\nHANDBOOK")
    r.bold = True
    r.font.size = Pt(34)
    r.font.color.rgb = ACCENT

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("A Practical Reference for Concepts, Systems,\nLoads and Design Behaviour")
    r.font.size = Pt(15)
    r.font.color.rgb = GREY

    doc.add_paragraph()
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run("Aligned with IS 456 : 2000  •  IS 875  •  IS 1893 : 2016  •  IS 13920 : 2016  •  IS 2911")
    r.font.size = Pt(11)
    r.font.color.rgb = ACCENT2
    r.italic = True

    for _ in range(8):
        doc.add_paragraph()
    d = doc.add_paragraph()
    d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = d.add_run("2026 Edition")
    r.font.size = Pt(12)
    r.font.color.rgb = GREY
    doc.add_page_break()


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HERE, "Structural_Engineering_Handbook.docx")

    files = sorted(
        f for f in os.listdir(HERE)
        if re.match(r"^\d{2}-.*\.md$", f)
    )
    if not files:
        print("No chapter markdown files found.")
        sys.exit(1)

    doc = Document()
    setup_styles(doc)

    # page margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    build_title_page(doc)

    for idx, fname in enumerate(files):
        with open(os.path.join(HERE, fname), encoding="utf-8") as fh:
            md = fh.read()
        # front matter (00) flows right after title page; chapters page-break
        process_markdown(doc, md, first_heading_page_break=(idx > 0))

    doc.save(out)
    size_kb = os.path.getsize(out) / 1024
    print(f"Written: {out} ({size_kb:.0f} KB) from {len(files)} markdown files")


if __name__ == "__main__":
    main()
